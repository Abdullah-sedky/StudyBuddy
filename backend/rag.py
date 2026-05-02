import glob
import os
import json
from dotenv import load_dotenv
from langchain_groq import ChatGroq
from langchain_community.document_loaders import (
    PyPDFLoader,
    DirectoryLoader,
)
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_chroma import Chroma
from langchain_core.documents import Document
from langchain_core.document_loaders import BaseLoader
from langchain_core.prompts import ChatPromptTemplate, PromptTemplate
from langchain_classic.chains.combine_documents import create_stuff_documents_chain
from langchain_classic.chains import create_retrieval_chain
from langchain_core.runnables.history import RunnableWithMessageHistory
from langchain_community.chat_message_histories import ChatMessageHistory

load_dotenv()

llm = ChatGroq(model_name="llama-3.3-70b-versatile")

system_prompt = """
    You are a helpful university tutor.

    1. Try to answer the user's question using ONLY the provided Context.
    2. If the answer is not in the Context, you may use your general knowledge, but you MUST start your answer by saying: 'Note: I am answering this based on general knowledge.'
    3. At the end of every response, create a 'Resources' section.
       - If you used the Context, list the document names (metadata source).
       - If you used general knowledge, list the specific external concepts or standard references you relied on.

    DIAGRAMS
    --------
    When the user asks you to explain something visually, draw a diagram, or show how something works,
    respond with a Mermaid diagram inside a fenced code block, like this:

    ```mermaid
    flowchart TD
        A[Start] --> B[Step]
        B --> C[End]
    ```

    Use the most appropriate diagram type for the concept:
    - flowchart TD / LR  — for processes, pipelines, algorithms
    - sequenceDiagram    — for interactions between components or actors
    - classDiagram       — for object relationships or data structures
    - mindmap            — for concept overviews and topic breakdowns

    Keep diagrams concise (under 20 nodes). You may combine a short text explanation with the diagram.
    Always use valid Mermaid syntax. Inside node labels, never use parentheses (), 
    brackets [], or any special characters. Use plain text only inside node labels.
    Context: {context}
"""

prompt = ChatPromptTemplate.from_messages([
    ("system", system_prompt),
    ("placeholder", "{chat_history}"),
    ("human", "{input}"),
])

document_prompt = PromptTemplate(
    input_variables=["page_content", "source"],
    template="Source File: {source}\nContent: {page_content}"
)

# Memory

store = {}  # session_id -> ChatMessageHistory, lives in RAM

def get_session_history(session_id: str) -> ChatMessageHistory:
    if session_id not in store:
        store[session_id] = ChatMessageHistory()
    return store[session_id]


def _iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if getattr(shape, "shapes", None):
            yield from _iter_shapes(shape.shapes)


def _safe_text(v) -> str:
    return (v or "").strip()


def _chart_to_lines(shape) -> list[str]:
    lines: list[str] = []
    chart = getattr(shape, "chart", None)
    if chart is None:
        return lines

    chart_title = ""
    try:
        if chart.has_title and chart.chart_title and chart.chart_title.text_frame:
            chart_title = _safe_text(chart.chart_title.text_frame.text)
    except Exception:
        chart_title = ""

    if chart_title:
        lines.append(f"Chart title: {chart_title}")
    else:
        lines.append("Chart present")

    try:
        categories = []
        if chart.plots and chart.plots[0].categories:
            categories = [_safe_text(c.label) for c in chart.plots[0].categories]
        for series in chart.series:
            values = []
            for idx, point in enumerate(series.points):
                cat = categories[idx] if idx < len(categories) else f"item-{idx + 1}"
                values.append(f"{cat}={point.value}")
            if values:
                lines.append(f"Series '{series.name}': " + ", ".join(values))
    except Exception:
        lines.append("Chart data could not be fully read.")

    return lines


def _picture_to_lines(shape) -> list[str]:
    lines: list[str] = []
    name = _safe_text(getattr(shape, "name", ""))
    alt_text = _safe_text(getattr(shape, "alternative_text", ""))
    if name or alt_text:
        lines.append(
            "Image: "
            + ", ".join(
                [part for part in [f"name={name}" if name else "", f"description={alt_text}" if alt_text else ""] if part]
            )
        )
    else:
        lines.append("Image present")
    return lines


def _text_from_slide(slide) -> list[str]:
    """Extract text and structured hints from a slide, including charts/images."""
    lines: list[str] = []

    notes_text = ""
    try:
        if slide.has_notes_slide and slide.notes_slide and slide.notes_slide.notes_text_frame:
            notes_text = _safe_text(slide.notes_slide.notes_text_frame.text)
    except Exception:
        notes_text = ""
    if notes_text:
        lines.append(f"Speaker notes: {notes_text}")

    for shape in _iter_shapes(slide.shapes):
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                for cell in row.cells:
                    t = _safe_text(cell.text)
                    if t:
                        lines.append(t)
            continue
        if getattr(shape, "has_chart", False):
            lines.extend(_chart_to_lines(shape))
            continue
        if getattr(shape, "image", None) is not None:
            lines.extend(_picture_to_lines(shape))
            continue
        if getattr(shape, "has_text_frame", False):
            for para in shape.text_frame.paragraphs:
                t = _safe_text(para.text)
                if t:
                    lines.append(t)
    return lines


class PptxTextLoader(BaseLoader):
    """Lightweight PPTX loader (no unstructured dependency)."""

    def __init__(self, file_path: str):
        self.file_path = os.path.normpath(file_path)

    def load(self) -> list[Document]:
        try:
            from pptx import Presentation
        except ImportError as e:
            raise ImportError(
                "PowerPoint support requires the 'python-pptx' package "
                "(import name 'pptx'). Install with: pip install python-pptx "
                "— use the same Python you use to run uvicorn."
            ) from e

        prs = Presentation(self.file_path)
        parts: list[str] = []
        for idx, slide in enumerate(prs.slides, start=1):
            slide_lines = _text_from_slide(slide)
            if slide_lines:
                parts.append(f"Slide {idx}:\n" + "\n".join(slide_lines))
        text = "\n\n".join(parts).strip()
        return [
            Document(
                page_content=text,
                metadata={"source": os.path.abspath(self.file_path)},
            )
        ]


# create vector space

current_vectorstore = None


def release_study_brain():
    """Best-effort close of the persisted Chroma client (helps unlock files on Windows)."""
    global current_vectorstore
    if current_vectorstore is None:
        return
    try:
        client = getattr(current_vectorstore, "_client", None)
        if client is not None:
            system = getattr(client, "_system", None)
            if system is not None and hasattr(system, "stop"):
                system.stop()
    except Exception as e:
        print(f"Chroma release (non-fatal): {e}")
    current_vectorstore = None


def create_study_brain(data_path="./data", persist_dir="./chroma_db"):
    global current_vectorstore

    release_study_brain()

    embeddings = HuggingFaceEmbeddings(model_name="all-MiniLM-L6-v2")

    if os.path.exists(persist_dir) and os.listdir(persist_dir):
        print("Loading existing vector store...")
        current_vectorstore = Chroma(
            persist_directory=persist_dir,
            embedding_function=embeddings
        )
        return current_vectorstore.as_retriever(search_kwargs={"k": 8})

    loaders = [
        DirectoryLoader(data_path, glob="**/*.pdf", loader_cls=PyPDFLoader, recursive=True),
        DirectoryLoader(data_path, glob="**/*.pptx", loader_cls=PptxTextLoader, recursive=True),
    ]

    pptx_on_disk = glob.glob(
        os.path.join(data_path, "**", "*.pptx"), recursive=True
    )
    pptx_on_disk += glob.glob(os.path.join(data_path, "*.pptx"), recursive=False)

    docs: list[Document] = []
    loader_errors: list[str] = []
    for loader in loaders:
        try:
            docs.extend(loader.load())
        except Exception as e:
            msg = f"{type(loader).__name__}: {e}"
            loader_errors.append(msg)
            print(f"Loader failed: {msg}")

    if not docs:
        hint = ""
        if loader_errors:
            hint = " Errors: " + "; ".join(loader_errors)
        raise ValueError("No documents found in data folder." + hint)

    if pptx_on_disk:
        pptx_sources = {
            os.path.normcase(os.path.abspath(p)) for p in pptx_on_disk
        }
        indexed_pptx = {
            os.path.normcase(os.path.abspath(str(d.metadata.get("source", ""))))
            for d in docs
            if str(d.metadata.get("source", "")).lower().endswith(".pptx")
        }
        missing = pptx_sources - indexed_pptx
        if missing:
            raise ValueError(
                "PPTX files were not indexed (loader produced no documents for them): "
                + ", ".join(sorted(missing))
            )

    print(f"Loaded {len(docs)} documents, building vector store...")

    splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=100)
    splits = splitter.split_documents(docs)
    splits = [s for s in splits if s.page_content and s.page_content.strip()]
    if not splits:
        raise ValueError(
            "No text could be extracted for search (files may be image-only or encrypted)."
        )

    current_vectorstore = Chroma.from_documents(
        documents=splits,
        embedding=embeddings,
        persist_directory=persist_dir
    )

    return current_vectorstore.as_retriever(search_kwargs={"k": 8})

# ---------- Chain Builder ----------

def build_chain(retriever):
    """
    Builds a fresh conversational RAG chain around the given retriever.
    Call this once at startup, and again every time a new file is uploaded.
    """
    document_chain = create_stuff_documents_chain(llm, prompt, document_prompt=document_prompt)
    rag_chain = create_retrieval_chain(retriever, document_chain)

    return RunnableWithMessageHistory(
        rag_chain,
        get_session_history,
        input_messages_key="input",
        history_messages_key="chat_history",
    )


def _context_from_docs(docs: list[Document], max_chars: int = 12000) -> str:
    chunks: list[str] = []
    total = 0
    for d in docs:
        source = str(d.metadata.get("source", "unknown"))
        text = (d.page_content or "").strip()
        if not text:
            continue
        piece = f"Source: {source}\n{text}\n"
        if total + len(piece) > max_chars:
            remaining = max_chars - total
            if remaining > 200:
                piece = piece[:remaining]
                chunks.append(piece)
            break
        chunks.append(piece)
        total += len(piece)
    return "\n---\n".join(chunks)


def generate_assessment(
    retriever,
    *,
    assessment_type: str,
    num_questions: int = 8,
    topic: str = "",
) -> dict:
    q_count = max(3, min(num_questions, 20))
    kind = "model exam" if assessment_type == "exam" else "mcq quiz"
    retrieval_query = (
        f"Create a {kind} about: {topic}" if topic else f"Create a {kind} from uploaded course material"
    )
    docs = retriever.invoke(retrieval_query)
    context = _context_from_docs(docs)
    if not context:
        raise ValueError("No indexed context available to generate assessment.")

    prompt = f"""
You are generating an interactive {kind} for a university student.
Use ONLY the provided context.

Return STRICT JSON only (no markdown fences, no extra text) with this shape:
{{
  "title": "string",
  "type": "{assessment_type}",
  "questions": [
    {{
      "question": "string",
      "options": ["string", "string", "string", "string"],
      "correct_index": 0,
      "explanation": "string"
    }}
  ]
}}

Rules:
- Exactly {q_count} questions.
- Exactly 4 options per question.
- correct_index must be an integer from 0 to 3.
- Questions must be clear and unambiguous.
- Keep options realistic and not obviously wrong.
- For exam mode: increase difficulty and include conceptual + applied questions.
- Do not include any text outside JSON.

Context:
{context}
"""

    response = llm.invoke(prompt)
    content = (getattr(response, "content", "") or "").strip()
    try:
        payload = json.loads(content)
    except json.JSONDecodeError:
        start = content.find("{")
        end = content.rfind("}")
        if start == -1 or end == -1 or end <= start:
            raise ValueError("Model output was not valid JSON.")
        payload = json.loads(content[start : end + 1])

    questions = payload.get("questions", [])
    if not isinstance(questions, list) or not questions:
        raise ValueError("Assessment generation returned no questions.")

    normalized = []
    for q in questions[:q_count]:
        options = q.get("options", [])
        if not isinstance(options, list) or len(options) < 4:
            continue
        correct_index = q.get("correct_index", 0)
        if not isinstance(correct_index, int) or correct_index < 0 or correct_index > 3:
            correct_index = 0
        normalized.append(
            {
                "question": str(q.get("question", "")).strip(),
                "options": [str(o).strip() for o in options[:4]],
                "correct_index": correct_index,
                "explanation": str(q.get("explanation", "")).strip(),
            }
        )

    if len(normalized) < 3:
        raise ValueError("Assessment generation failed quality checks.")

    return {
        "title": str(payload.get("title", "Practice Assessment")).strip() or "Practice Assessment",
        "type": assessment_type,
        "questions": normalized,
    }